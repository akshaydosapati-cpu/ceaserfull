"""
File Parser for Ceaser Predictions
Handles parsing of various file formats
"""

import os
import logging
from typing import Any, Dict, Optional
import json

logger = logging.getLogger(__name__)

class FileParser:
    """Parse different file formats"""
    
    def __init__(self):
        self.parsers = {
            "pdf": self._parse_pdf,
            "docx": self._parse_docx,
            "txt": self._parse_txt,
            "xlsx": self._parse_excel,
            "xls": self._parse_excel,
            "csv": self._parse_csv,
            "json": self._parse_json,
            "xml": self._parse_xml,
            "jpg": self._parse_image,
            "jpeg": self._parse_image,
            "png": self._parse_image,
            "gif": self._parse_image,
        }
        self.stats = {"files_parsed": 0}
    
    async def initialize(self):
        """Initialize parser"""
        logger.info("Initializing File Parser...")
        logger.info("File Parser ready!")
    
    def get_stats(self):
        """Get parser statistics"""
        return self.stats
    
    def get_file_type(self, filename: str) -> str:
        """Get file extension"""
        return filename.split('.')[-1].lower() if '.' in filename else "unknown"
    
    def parse_content(self, content: bytes, filename: str) -> Any:
        """Parse file content based on file type"""
        try:
            file_type = self.get_file_type(filename)
            parser = self.parsers.get(file_type, self._parse_generic)
            
            logger.info(f"Parsing {file_type} file: {filename}")
            parsed = parser(content)
            
            self.stats["files_parsed"] += 1
            return parsed
            
        except Exception as e:
            logger.error(f"Error parsing file {filename}: {e}")
            return str(content[:1000])  # Return first 1000 bytes as text
    
    def _parse_pdf(self, content: bytes) -> str:
        """Parse PDF file"""
        try:
            import PyPDF2
            from io import BytesIO
            
            pdf_file = BytesIO(content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
            return text
        except ImportError:
            logger.warning("PyPDF2 not installed. PDF parsing limited.")
            return "PDF content (requires PyPDF2 for full parsing)"
        except Exception as e:
            logger.error(f"Error parsing PDF: {e}")
            return "PDF content parsing error"
    
    def _parse_docx(self, content: bytes) -> str:
        """Parse DOCX file"""
        try:
            from docx import Document
            from io import BytesIO
            
            doc = Document(BytesIO(content))
            text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            return text
        except ImportError:
            logger.warning("python-docx not installed. DOCX parsing limited.")
            return "DOCX content (requires python-docx for full parsing)"
        except Exception as e:
            logger.error(f"Error parsing DOCX: {e}")
            return "DOCX content parsing error"
    
    def _parse_txt(self, content: bytes) -> str:
        """Parse TXT file"""
        try:
            return content.decode('utf-8', errors='ignore')
        except Exception:
            return content.decode('latin-1', errors='ignore')
    
    def _parse_excel(self, content: bytes) -> Dict[str, Any]:
        """Parse Excel file"""
        try:
            from openpyxl import load_workbook
            from io import BytesIO

            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                return {"type": "excel", "rows": [], "summary": "0 rows, 0 columns"}
            headers = [str(value) if value is not None else f"column_{index + 1}" for index, value in enumerate(values[0])]
            rows = [dict(zip(headers, row)) for row in values[1:]]
            return {
                "type": "excel",
                "rows": rows,
                "summary": f"{len(rows)} rows, {len(headers)} columns"
            }
        except ImportError:
            logger.warning("openpyxl not installed. Excel parsing limited.")
            return {"type": "excel", "error": "Requires openpyxl"}
        except Exception as e:
            logger.error(f"Error parsing Excel: {e}")
            return {"type": "excel", "error": str(e)}
    
    def _parse_csv(self, content: bytes) -> str:
        """Parse CSV file"""
        try:
            return content.decode('utf-8', errors='ignore')
        except Exception:
            return content.decode('latin-1', errors='ignore')
    
    def _parse_json(self, content: bytes) -> Dict[str, Any]:
        """Parse JSON file"""
        try:
            return json.loads(content.decode('utf-8'))
        except Exception as e:
            logger.error(f"JSON parsing error: {e}")
            return {"error": "Invalid JSON"}
    
    def _parse_xml(self, content: bytes) -> str:
        """Parse XML file"""
        try:
            return content.decode('utf-8', errors='ignore')
        except Exception:
            return content.decode('latin-1', errors='ignore')
    
    def _parse_image(self, content: bytes) -> str:
        """Parse image file (OCR)"""
        try:
            from PIL import Image
            import pytesseract
            from io import BytesIO
            
            image = Image.open(BytesIO(content))
            text = pytesseract.image_to_string(image)
            return text
        except ImportError:
            logger.warning("pytesseract or Pillow not installed. Image OCR not available.")
            return "Image content (requires pytesseract and Pillow for OCR)"
        except Exception as e:
            logger.error(f"Error parsing image: {e}")
            return "Image OCR error"
    
    def _parse_generic(self, content: bytes) -> str:
        """Parse unknown file type"""
        try:
            # Try to decode as text
            return content.decode('utf-8', errors='ignore')[:5000]  # Limit to 5000 chars
        except Exception:
            return f"Binary content (size: {len(content)} bytes)"

